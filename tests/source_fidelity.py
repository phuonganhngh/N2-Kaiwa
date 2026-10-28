"""Compare every imported cell directly with the archived workbook."""
import hashlib
import json
from pathlib import Path
import openpyxl
root = Path(__file__).resolve().parents[1]
data = json.loads((root / 'data/grammar.json').read_text())
source = root / 'source/Grammar-N2.xlsx'
assert hashlib.sha256(source.read_bytes()).hexdigest() == data['sha256']
w = openpyxl.load_workbook(source, data_only=False)
assert w.sheetnames == ['Sheet']
sheet = w['Sheet']
assert list(next(sheet.values)) == data['headers']
entries, notes = [], []
for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
    if row[0] is not None:
        entries.append(dict(zip(['stt','structure','meaning','examples'],row),sourceRow=row_number))
    elif any(v is not None for v in row):
        notes.append(dict(sourceRow=row_number,afterStt=entries[-1]['stt'],cells=list(row)))
assert data['entries'] == entries
assert data['notes'] == notes
print(f'PASS: all four original columns match exactly for {len(entries)} entries; {len(notes)} notes preserved.')
